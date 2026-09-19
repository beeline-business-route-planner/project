from __future__ import annotations

from io import BytesIO
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook


def _import(client: TestClient, content: bytes) -> tuple[str, str]:
    response = client.post(
        "/api/v1/imports",
        files={
            "file": (
                "Восток Синтетические данные.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers={"Idempotency-Key": "fixture-east-v1"},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    return body["dataset_id"], body["scenario_id"]


def test_invalid_date_and_window_return_validation_errors(client: TestClient) -> None:
    invalid_date = client.get(
        "/api/v1/requests",
        params={"scenario_id": str(uuid4()), "planning_date": "not-a-date"},
    )
    assert invalid_date.status_code == 422, invalid_date.text
    assert invalid_date.json()["code"] == "validation_error"

    reversed_window = client.post(
        "/api/v1/requests",
        json={
            "scenario_id": str(uuid4()),
            "planning_date": "2026-08-17",
            "external_id": "invalid-window",
            "address": "Москва, тестовый адрес",
            "district": "Тестовый",
            "window_start": "2026-08-17T12:00:00+03:00",
            "window_end": "2026-08-17T10:00:00+03:00",
            "service_minutes": 30,
            "required_skill": "connection",
            "priority": "normal",
            "idempotency_key": "invalid-window",
            "actor": "test",
        },
    )
    assert reversed_window.status_code == 422, reversed_window.text
    assert reversed_window.json()["code"] == "validation_error"


def test_import_plan_approve_replan_diff_and_reports(
    client: TestClient, synthetic_xlsx: bytes
) -> None:
    _, scenario_id = _import(client, synthetic_xlsx)

    duplicate = client.post(
        "/api/v1/imports",
        files={"file": ("Восток Синтетические данные.xlsx", synthetic_xlsx)},
        headers={"Idempotency-Key": "fixture-east-v1"},
    )
    assert duplicate.status_code == 200
    assert duplicate.json()["duplicate"] is True

    requests = client.get(
        "/api/v1/requests",
        params={"scenario_id": scenario_id, "planning_date": "2026-08-17"},
    )
    assert requests.status_code == 200
    assert len(requests.json()) == 2

    run = client.post(
        "/api/v1/plans/run",
        json={"scenario_id": scenario_id, "planning_date": "2026-08-17"},
    )
    assert run.status_code == 200, run.text
    first_plan = run.json()["plan_id"]
    plan = client.get(f"/api/v1/plans/{first_plan}")
    assert plan.status_code == 200
    assert len(plan.json()["assignments"]) == 2
    assert plan.json()["metrics"]["travel_time_seconds"] >= 0

    approved = client.post(
        f"/api/v1/plans/{first_plan}/approve",
        json={"expected_base_plan_id": None, "actor": "dispatcher"},
    )
    assert approved.status_code == 200, approved.text

    event = client.post(
        "/api/v1/events",
        json={
            "scenario_id": scenario_id,
            "planning_date": "2026-08-17",
            "event_type": "fact_changed",
            "effective_at": "2026-08-17T10:00:00+03:00",
            "payload": {"note": "demo"},
            "idempotency_key": "event-1",
            "actor": "dispatcher",
        },
    )
    assert event.status_code == 200, event.text
    event_id = event.json()["event_id"]
    second_plan = event.json()["replanning"]["plan_id"]
    assert second_plan != first_plan

    audit = client.get("/api/v1/audit", params={"scenario_id": scenario_id})
    assert audit.status_code == 200, audit.text
    day_event_audit = next(
        item for item in audit.json() if item["action"] == "day_event_created"
    )
    assert day_event_audit["object_id"] == event_id

    diff = client.get(
        "/api/v1/plans/diff",
        params={"old_plan_id": first_plan, "new_plan_id": second_plan},
    )
    assert diff.status_code == 200, diff.text
    assert diff.json()["summary"]["unchanged"] == 2

    stale = client.post(
        f"/api/v1/plans/{second_plan}/approve",
        json={"expected_base_plan_id": None, "actor": "dispatcher"},
    )
    assert stale.status_code == 409

    approved_second = client.post(
        f"/api/v1/plans/{second_plan}/approve",
        json={"expected_base_plan_id": first_plan, "actor": "dispatcher"},
    )
    assert approved_second.status_code == 200, approved_second.text

    for report_format, signature in (("xlsx", b"PK"), ("pdf", b"%PDF")):
        report = client.get(f"/api/v1/plans/{second_plan}/reports/{report_format}")
        assert report.status_code == 200, report.text
        assert report.content.startswith(signature)

    dashboard = client.get(
        "/api/v1/dashboard",
        params={"scenario_id": scenario_id, "planning_date": "2026-08-17"},
    )
    assert dashboard.status_code == 200
    assert dashboard.json()["active_plan"]["id"] == second_plan


def test_planned_finish_does_not_create_completed_fact(
    client: TestClient, synthetic_xlsx: bytes
) -> None:
    _, scenario_id = _import(client, synthetic_xlsx)
    response = client.get(
        "/api/v1/requests",
        params={"scenario_id": scenario_id, "planning_date": "2026-08-17"},
    )
    assert {item["status"] for item in response.json()} == {"NOT_SENT"}


def test_report_preserves_actual_start_and_tracks_fact_version(
    client: TestClient, synthetic_xlsx: bytes
) -> None:
    _, scenario_id = _import(client, synthetic_xlsx)
    run = client.post(
        "/api/v1/plans/run",
        json={"scenario_id": scenario_id, "planning_date": "2026-08-17"},
    )
    assert run.status_code == 200, run.text
    plan_id = run.json()["plan_id"]
    assignment = client.get(f"/api/v1/plans/{plan_id}").json()["assignments"][0]

    initial_report = client.get(f"/api/v1/plans/{plan_id}/reports/xlsx")
    assert initial_report.status_code == 200, initial_report.text
    initial_sheet = load_workbook(BytesIO(initial_report.content)).active
    initial_version = initial_sheet.cell(row=3, column=2).value

    facts = (
        ("SENT", "2026-08-17T08:05:00+03:00", {}),
        ("EN_ROUTE", "2026-08-17T08:10:00+03:00", {}),
        (
            "IN_PROGRESS",
            "2026-08-17T10:15:00+03:00",
            {"actual_start": "2026-08-17T10:15:00+03:00"},
        ),
        (
            "COMPLETED",
            "2026-08-17T11:00:00+03:00",
            {"actual_finish": "2026-08-17T11:00:00+03:00"},
        ),
    )
    for status, effective_at, actual_times in facts:
        response = client.post(
            f"/api/v1/requests/{assignment['request_id']}/facts",
            json={
                "status": status,
                "effective_at": effective_at,
                "actor": "dispatcher",
                "reason": "report regression",
                **actual_times,
            },
        )
        assert response.status_code == 200, response.text

    updated_report = client.get(f"/api/v1/plans/{plan_id}/reports/xlsx")
    assert updated_report.status_code == 200, updated_report.text
    updated_sheet = load_workbook(BytesIO(updated_report.content)).active
    assert updated_sheet.cell(row=3, column=2).value != initial_version
    report_row = next(
        row
        for row in updated_sheet.iter_rows(min_row=7, values_only=True)
        if str(row[2]) == assignment["request_external_id"]
    )
    assert report_row[6] == "COMPLETED"
    assert report_row[7] is not None
    assert report_row[8] is not None
    assert report_row[8] > report_row[7]


def test_manual_change_recalculates_and_respects_skill_and_execution_lock(
    client: TestClient, synthetic_xlsx: bytes
) -> None:
    _, scenario_id = _import(client, synthetic_xlsx)
    run = client.post(
        "/api/v1/plans/run",
        json={"scenario_id": scenario_id, "planning_date": "2026-08-17"},
    ).json()
    plan_id = run["plan_id"]
    assert client.post(
        f"/api/v1/plans/{plan_id}/approve",
        json={"expected_base_plan_id": None, "actor": "dispatcher"},
    ).status_code == 200
    plan = client.get(f"/api/v1/plans/{plan_id}").json()
    assignment = plan["assignments"][0]
    engineers = client.get("/api/v1/engineers", params={"scenario_id": scenario_id}).json()
    incompatible = next(item for item in engineers if "local" not in item["skills"])
    compatible = next(
        item
        for item in engineers
        if "local" in item["skills"] and item["id"] != assignment["engineer_id"]
    )

    invalid = client.post(
        f"/api/v1/plans/{plan_id}/manual-change",
        json={
            "request_id": assignment["request_id"],
            "engineer_id": incompatible["id"],
            "position": 1,
            "reason": "skill validation",
            "actor": "dispatcher",
        },
    )
    assert invalid.status_code == 422
    assert invalid.json()["code"] == "missing_skill"

    changed = client.post(
        f"/api/v1/plans/{plan_id}/manual-change",
        json={
            "request_id": assignment["request_id"],
            "engineer_id": compatible["id"],
            "position": 1,
            "reason": "dispatcher choice",
            "actor": "dispatcher",
        },
    )
    assert changed.status_code == 200, changed.text
    changed_plan_id = changed.json()["plan_id"]
    old_after = client.get(f"/api/v1/plans/{plan_id}").json()
    new_after = client.get(f"/api/v1/plans/{changed_plan_id}").json()
    old_engineer = next(
        item["engineer_id"]
        for item in old_after["assignments"]
        if item["request_id"] == assignment["request_id"]
    )
    new_engineer = next(
        item["engineer_id"]
        for item in new_after["assignments"]
        if item["request_id"] == assignment["request_id"]
    )
    assert old_engineer == assignment["engineer_id"]
    assert new_engineer == compatible["id"]

    for status, effective_at in (
        ("SENT", "2026-08-17T08:05:00+03:00"),
        ("EN_ROUTE", "2026-08-17T08:10:00+03:00"),
    ):
        fact = client.post(
            f"/api/v1/requests/{assignment['request_id']}/facts",
            json={
                "status": status,
                "effective_at": effective_at,
                "actor": "dispatcher",
                "reason": "demo",
            },
        )
        assert fact.status_code == 200, fact.text
    locked = client.post(
        f"/api/v1/plans/{plan_id}/manual-change",
        json={
            "request_id": assignment["request_id"],
            "engineer_id": compatible["id"],
            "position": 1,
            "reason": "must fail",
            "actor": "dispatcher",
        },
    )
    assert locked.status_code == 409
    assert locked.json()["code"] == "request_locked"


def test_only_one_candidate_can_win_concurrent_approval(
    client: TestClient, synthetic_xlsx: bytes
) -> None:
    _, scenario_id = _import(client, synthetic_xlsx)
    first_plan = client.post(
        "/api/v1/plans/run",
        json={"scenario_id": scenario_id, "planning_date": "2026-08-17"},
    ).json()["plan_id"]
    assert client.post(
        f"/api/v1/plans/{first_plan}/approve",
        json={"expected_base_plan_id": None, "actor": "dispatcher"},
    ).status_code == 200

    drafts: list[str] = []
    for _ in (1, 2):
        response = client.post(
            "/api/v1/plans/run",
            json={
                "scenario_id": scenario_id,
                "planning_date": "2026-08-17",
                "base_plan_id": first_plan,
            },
        )
        assert response.status_code == 200, response.text
        drafts.append(response.json()["plan_id"])
    winner = client.post(
        f"/api/v1/plans/{drafts[0]}/approve",
        json={"expected_base_plan_id": first_plan, "actor": "dispatcher"},
    )
    assert winner.status_code == 200, winner.text
    loser = client.post(
        f"/api/v1/plans/{drafts[1]}/approve",
        json={"expected_base_plan_id": first_plan, "actor": "dispatcher"},
    )
    assert loser.status_code == 409
    assert loser.json()["code"] == "stale_base_plan"


def test_cancellation_event_removes_request_from_active_routes(
    client: TestClient, synthetic_xlsx: bytes
) -> None:
    _, scenario_id = _import(client, synthetic_xlsx)
    requests = client.get(
        "/api/v1/requests",
        params={"scenario_id": scenario_id, "planning_date": "2026-08-17"},
    ).json()
    cancelled_request_id = requests[0]["id"]

    event = client.post(
        "/api/v1/events",
        json={
            "scenario_id": scenario_id,
            "planning_date": "2026-08-17",
            "event_type": "request_cancelled",
            "effective_at": "2026-08-17T09:00:00+03:00",
            "payload": {"request_id": cancelled_request_id},
            "idempotency_key": "cancel-request-1",
            "actor": "dispatcher",
        },
    )
    assert event.status_code == 200, event.text
    replanned = client.get(f"/api/v1/plans/{event.json()['replanning']['plan_id']}")
    assert replanned.status_code == 200, replanned.text
    plan = replanned.json()
    assert cancelled_request_id not in {
        assignment["request_id"] for assignment in plan["assignments"]
    }
    assert {
        item["reason_code"]
        for item in plan["unassigned"]
        if item["request_id"] == cancelled_request_id
    } == {"cancelled"}

    current_requests = client.get(
        "/api/v1/requests",
        params={"scenario_id": scenario_id, "planning_date": "2026-08-17"},
    ).json()
    assert next(
        item["status"] for item in current_requests if item["id"] == cancelled_request_id
    ) == "CANCELLED"

    duplicate = client.post(
        "/api/v1/events",
        json={
            "scenario_id": scenario_id,
            "planning_date": "2026-08-17",
            "event_type": "request_cancelled",
            "effective_at": "2026-08-17T09:00:00+03:00",
            "payload": {"request_id": cancelled_request_id},
            "idempotency_key": "cancel-request-1",
            "actor": "dispatcher",
        },
    )
    assert duplicate.status_code == 200, duplicate.text
    assert duplicate.json()["duplicate"] is True
    assert duplicate.json()["replanning"] is None

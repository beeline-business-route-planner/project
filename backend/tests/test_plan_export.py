import io
import unittest
import uuid
from dataclasses import replace
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from unittest.mock import AsyncMock

from openpyxl import load_workbook

from src.api.plans.diff_dto import (
    PlanMetricsDTO,
    PlanSnapshotDTO,
    SnapshotEngineerDTO,
    SnapshotRequestDTO,
)
from src.api.plans.export import PlanXlsxExporter
from src.api.exc.plans import PlanExportStorageError
from src.api.plans.service import PlanExportService
from src.core.s3 import ExportDownload, ExportKind, S3UnavailableError
from src.core.db.enums import (
    ApprovalStatus,
    PlanKind,
    Region,
    RequestTypeBk,
    RequestTypeHd,
    Skill,
    UnassignedReason,
    VehicleType,
)


class PlanXlsxExporterTest(unittest.TestCase):
    def setUp(self) -> None:
        start = datetime(2026, 9, 23, 10)
        assigned = SnapshotRequestDTO(
            request_id=uuid.UUID(int=11),
            external_id=11,
            address='=HYPERLINK("unsafe")',
            district="Район",
            latitude=Decimal("55.750000"),
            longitude=Decimal("37.610000"),
            window_start=start,
            window_end=start + timedelta(hours=2),
            priority=1,
            required_skill=Skill.CONNECTION_AND_ORDERS,
            service_minutes=60,
            engineer_id=uuid.UUID(int=21),
            sequence_number=1,
            planned_arrival=start,
            planned_start=start,
            planned_finish=start + timedelta(hours=1),
            travel_minutes=10,
            distance_km=Decimal("4.25"),
            is_locked=False,
            unassigned_reason=None,
            upload_id=uuid.UUID(int=99),
            type_bk=RequestTypeBk.CONNECTION,
            type_hd=RequestTypeHd.CONNECTION_REQUEST,
            is_gigabit=False,
            norm_minutes=80,
        )
        urgent = replace(
            assigned,
            request_id=uuid.UUID(int=12),
            external_id=12,
            address="Срочная заявка",
            engineer_id=None,
            sequence_number=None,
            planned_arrival=None,
            planned_start=None,
            planned_finish=None,
            travel_minutes=None,
            distance_km=None,
            unassigned_reason=UnassignedReason.NO_AVAILABLE_ENGINEER,
            upload_id=None,
        )
        engineer = SnapshotEngineerDTO(
            engineer_id=uuid.UUID(int=21),
            name="Инженер / с очень длинным одинаковым именем",
            vehicle_type=VehicleType.CAR,
            shift_start=start,
            shift_end=start + timedelta(hours=8),
            start_latitude=Decimal("55.750000"),
            start_longitude=Decimal("37.610000"),
            is_available=True,
            route_distance_km=Decimal("4.25"),
            workload_without_travel=Decimal("12.50"),
            workload_with_travel=Decimal("14.58"),
            requests=(assigned,),
        )
        unused = replace(
            engineer,
            engineer_id=uuid.UUID(int=22),
            requests=(),
            route_distance_km=Decimal("0"),
            workload_without_travel=Decimal("0"),
            workload_with_travel=Decimal("0"),
        )
        self.snapshot = PlanSnapshotDTO(
            id=uuid.UUID(int=1),
            region=Region.VOSTOK,
            planning_date=date(2026, 9, 23),
            kind=PlanKind.EVENT_REPLAN,
            approval_status=ApprovalStatus.PENDING,
            created_at=datetime(2026, 9, 23, 7),
            approved_at=None,
            rejected_at=None,
            calculation_cutoff_at=start,
            based_on_plan_id=uuid.UUID(int=2),
            triggered_by_event_id=uuid.UUID(int=3),
            metrics=PlanMetricsDTO(
                assigned_requests_count=1,
                unassigned_requests_count=1,
                engineers_used_count=1,
                available_engineers_count=2,
                total_mileage_km=Decimal("4.25"),
                total_work_minutes=60,
                total_travel_minutes=10,
                average_workload_without_travel=Decimal("6.25"),
                average_workload_with_travel=Decimal("7.29"),
                average_used_workload_without_travel=Decimal("12.50"),
                average_used_workload_with_travel=Decimal("14.58"),
                min_workload_with_travel=Decimal("0"),
                max_workload_with_travel=Decimal("14.58"),
            ),
            requests=(urgent, assigned),
            engineers=(unused, engineer),
        )

    def test_all_plan_kinds_and_statuses_open_with_required_sheets(self) -> None:
        for kind in PlanKind:
            for status in ApprovalStatus:
                with self.subTest(kind=kind, status=status):
                    snapshot = replace(self.snapshot, kind=kind, approval_status=status)
                    workbook = load_workbook(io.BytesIO(PlanXlsxExporter.build(snapshot)))
                    self.assertEqual(
                        workbook.sheetnames[:5],
                        ["План", "Заявки", "Инженеры", "Неназначенные", "Метрики"],
                    )
                    self.assertEqual(len(workbook.sheetnames), 7)
                    self.assertEqual(len(set(name.casefold() for name in workbook.sheetnames)), 7)
                    self.assertTrue(all(len(name) <= 31 for name in workbook.sheetnames))
                    self.assertEqual(workbook["План"]["B5"].value, kind.value)
                    self.assertEqual(workbook["План"]["B6"].value, status.value)

    def test_rows_metrics_urgent_request_and_formula_safety(self) -> None:
        workbook = load_workbook(io.BytesIO(PlanXlsxExporter.build(self.snapshot)))
        requests = workbook["Заявки"]
        self.assertEqual([requests.cell(row, 2).value for row in (2, 3)], [11, 12])
        self.assertEqual(requests["C2"].data_type, "s")
        self.assertEqual(requests["L3"].value, "не назначена")
        self.assertIsNone(requests["U3"].value)
        self.assertEqual(requests["V2"].value, RequestTypeBk.CONNECTION.value)
        self.assertEqual(workbook["Неназначенные"]["B2"].value, 12)
        engineers = workbook["Инженеры"]
        self.assertEqual(sum(engineers.cell(row, 10).value for row in (2, 3)), 60)
        self.assertEqual(sum(engineers.cell(row, 11).value for row in (2, 3)), 10)
        self.assertEqual(sum(engineers.cell(row, 12).value for row in (2, 3)), 4.25)
        self.assertEqual(workbook["План"]["B7"].value, datetime(2026, 9, 23, 10))
        self.assertEqual(workbook.worksheets[5]["A4"].value, 1)

    def test_zero_assigned_stops_is_valid(self) -> None:
        unassigned = replace(
            self.snapshot.requests[1],
            engineer_id=None,
            sequence_number=None,
            planned_arrival=None,
            planned_start=None,
            planned_finish=None,
            travel_minutes=None,
            distance_km=None,
            unassigned_reason=UnassignedReason.NO_TIME_SLOT,
        )
        snapshot = replace(
            self.snapshot,
            requests=(self.snapshot.requests[0], unassigned),
            engineers=tuple(replace(item, requests=()) for item in self.snapshot.engineers),
        )
        workbook = load_workbook(io.BytesIO(PlanXlsxExporter.build(snapshot)))
        self.assertEqual(workbook["Заявки"].max_row, 3)
        self.assertEqual(workbook.worksheets[5].max_row, 3)


class PlanExportServiceTest(unittest.IsolatedAsyncioTestCase):
    async def test_exports_exact_snapshot_without_writing_database(self) -> None:
        snapshot = PlanXlsxExporterTest()
        snapshot.setUp()
        plans = AsyncMock()
        plans.get_snapshot_by_id.return_value = snapshot.snapshot
        delivery = AsyncMock()
        delivery.deliver.return_value = ExportDownload(
            url="https://example.invalid/export",
            expires_at=datetime(2026, 9, 23, 10, tzinfo=UTC),
            filename=f"plan-{snapshot.snapshot.id}.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size_bytes=123,
        )

        result = await PlanExportService(plans, delivery).export(snapshot.snapshot.id)

        self.assertEqual(result.url, "https://example.invalid/export")
        self.assertIsNotNone(result.expires_at.tzinfo)
        plans.get_snapshot_by_id.assert_awaited_once_with(snapshot.snapshot.id)
        arguments = delivery.deliver.await_args.kwargs
        self.assertEqual(arguments["kind"], ExportKind.PLAN)
        self.assertEqual(arguments["planning_date"], snapshot.snapshot.planning_date)
        self.assertEqual(arguments["filename"], f"plan-{snapshot.snapshot.id}.xlsx")
        workbook = load_workbook(io.BytesIO(arguments["data"]))
        self.assertEqual(workbook["План"]["B2"].value, str(snapshot.snapshot.id))

    async def test_storage_failure_is_controlled(self) -> None:
        snapshot = PlanXlsxExporterTest()
        snapshot.setUp()
        plans = AsyncMock()
        plans.get_snapshot_by_id.return_value = snapshot.snapshot
        delivery = AsyncMock()
        delivery.deliver.side_effect = S3UnavailableError("unavailable")
        with self.assertRaises(PlanExportStorageError):
            await PlanExportService(plans, delivery).export(snapshot.snapshot.id)


if __name__ == "__main__":
    unittest.main()

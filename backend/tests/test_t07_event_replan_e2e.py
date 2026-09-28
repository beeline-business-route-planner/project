"""T07 lifecycle against a disposable migrated local t07_test_* PostgreSQL database."""

import os
import unittest
import uuid
from datetime import datetime, timedelta
from io import BytesIO
from unittest.mock import patch
from zoneinfo import ZoneInfo

from fastapi import UploadFile
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.engine import make_url
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from src.api.exc.planning import (
    PlanningCurrentPlanMissing,
    PlanningEngineerStateConflict,
    PlanningEventTargetMissing,
    PlanningPendingEventExists,
    PlanningRequestAlreadyCancelled,
    PlanningRoutingUnavailable,
    PlanningUrgentRequestExists,
    PlanningUrgentRequestInvalid,
)
from src.api.exc.plans import PlanNotPendingError, PlanStateChangedError
from src.api.planning.router import create_event, import_initial_planning_data
from src.api.planning.schemas import EventPlanningRequest, UrgentRequestPayload
from src.api.planning.service import PlanningService
from src.api.plans.diff import PlanDiffEngine
from src.api.plans.export import PlanXlsxExporter
from src.api.plans.service import PlanService
from src.api.reports.service import DailyReportService
from src.core.algorithm import AlgorithmService
from src.core.algorithm.exc import AlgorithmAuditError
from src.core.db.enums import (
    ApprovalStatus,
    PlanKind,
    Region,
    ReplanningEventType,
    RequestStatus,
    RequestTypeBk,
    RequestTypeHd,
    Skill,
)
from src.core.db.models import Engineer, Plan, ReplanningEvent, Request
from src.core.db.uow import UnitOfWork
from tests.test_g02_initial_e2e import FakeGeocoder, FakeMatrixService, FakeStorage, pair

DATABASE_URL = os.environ.get("T07_TEST_DATABASE_URL")


@unittest.skipUnless(DATABASE_URL, "T07_TEST_DATABASE_URL is not set")
class EventReplanEndToEndTest(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self) -> None:
        assert DATABASE_URL is not None
        url = make_url(DATABASE_URL)
        if (
            url.drivername != "postgresql+asyncpg"
            or url.host not in {"127.0.0.1", "localhost", "postgres"}
            or not (url.database or "").startswith("t07_test_")
        ):
            self.fail("T07 requires a disposable local t07_test_* database")
        self.engine = create_async_engine(
            DATABASE_URL, connect_args={"server_settings": {"timezone": "UTC"}}
        )
        self.sessions = async_sessionmaker(self.engine, expire_on_commit=False)
        self.storage = FakeStorage()
        async with self.sessions() as session:
            if await session.scalar(select(func.count()).select_from(Plan)):
                self.fail("T07 database must be freshly migrated and empty")

    async def asyncTearDown(self) -> None:
        await self.engine.dispose()

    def planning(self, session, matrix=None) -> PlanningService:
        return PlanningService(
            UnitOfWork(session),
            FakeGeocoder(),
            self.storage,
            AlgorithmService(),
            matrix or FakeMatrixService(),
        )

    async def test_four_transitions_and_atomic_rollback(self) -> None:
        now = datetime.now(ZoneInfo("Europe/Moscow")).replace(tzinfo=None)
        if now.hour == 23:
            self.skipTest("Initial approval needs enough of the working day left")
        files = pair("Восток")
        uploads = [UploadFile(file=BytesIO(item.data), filename=item.filename) for item in files]
        async with self.sessions() as session:
            initial = await import_initial_planning_data(self.planning(session), uploads)
        self.assertEqual(initial.status, "success")
        base_id = initial.regions[0].plan_summary.id
        async with self.sessions() as session:
            await PlanService(UnitOfWork(session)).approve(base_id)
            original = (
                await session.scalars(select(Request).where(Request.upload_id.is_not(None)))
            ).one()
            engineer = (await session.scalars(select(Engineer))).one()
            request_id, engineer_id = original.id, engineer.id

        urgent = EventPlanningRequest(
            region=Region.VOSTOK,
            event_type=ReplanningEventType.URGENT_REQUEST,
            urgent_request=UrgentRequestPayload(
                external_id=987654321,
                type_bk=RequestTypeBk.GLOBAL_PROBLEM,
                type_hd=RequestTypeHd.EMERGENCY,
                district="Тестовый район",
                address="Тестовый адрес",
                is_gigabit=False,
                window_start=now + timedelta(minutes=20),
                window_end=datetime.combine(now.date(), datetime.max.time()).replace(microsecond=0),
                norm_minutes=100,
                norm_minutes_without_travel=80,
                priority=1,
                required_skill=Skill.EMERGENCY_WORKS,
            ),
        )
        async with self.sessions() as session:
            invalid = EventPlanningRequest(
                region=Region.VOSTOK,
                event_type=ReplanningEventType.URGENT_REQUEST,
                urgent_request=UrgentRequestPayload.model_validate(
                    {**urgent.urgent_request.model_dump(), "priority": 2}
                ),
            )
            with self.assertRaises(PlanningUrgentRequestInvalid):
                await create_event(invalid, self.planning(session))
        async with self.sessions() as session:
            with self.assertRaises(PlanningRoutingUnavailable):
                await create_event(
                    urgent, self.planning(session, FakeMatrixService(fail_on_call=1))
                )
        async with self.sessions() as session:
            self.assertEqual(
                await session.scalar(select(func.count()).select_from(ReplanningEvent)), 0
            )
            self.assertEqual(await session.scalar(select(func.count()).select_from(Plan)), 1)
            self.assertIsNone(
                (
                    await session.scalars(select(Request).where(Request.external_id == 987654321))
                ).first()
            )

        async with self.sessions() as session:
            with patch.object(
                AlgorithmService, "plan_replan", side_effect=AlgorithmAuditError("audit failed")
            ):
                with self.assertRaises(AlgorithmAuditError):
                    await create_event(urgent, self.planning(session))
        async with self.sessions() as session:
            with patch.object(UnitOfWork, "commit", side_effect=RuntimeError("database failed")):
                with self.assertRaises(RuntimeError):
                    await create_event(urgent, self.planning(session))
        async with self.sessions() as session:
            self.assertEqual(
                await session.scalar(select(func.count()).select_from(ReplanningEvent)), 0
            )
            self.assertEqual(await session.scalar(select(func.count()).select_from(Plan)), 1)
            self.assertIsNone(
                (
                    await session.scalars(select(Request).where(Request.external_id == 987654321))
                ).first()
            )

        async with self.sessions() as session:
            pending = await create_event(urgent, self.planning(session))
        self.assertEqual(pending.plan.kind, PlanKind.EVENT_REPLAN)
        self.assertEqual(pending.plan.based_on_plan_id, base_id)
        async with self.sessions() as session:
            with self.assertRaises(PlanningPendingEventExists):
                await create_event(urgent, self.planning(session))
        async with self.sessions() as session:
            decisions = PlanService(UnitOfWork(session))
            detail = await decisions.get_by_id(pending.plan.id)
            self.assertTrue(
                any(
                    item.request_id == pending.request_id
                    for group in detail.request_groups
                    for item in group.requests
                )
            )
            candidate = await decisions.get_snapshot_by_id(pending.plan.id)
            base = await decisions.get_snapshot_by_id(base_id)
            diff = PlanDiffEngine.compare(base, candidate)
            self.assertTrue(any(item.request_id == pending.request_id for item in diff.requests))
            book = load_workbook(BytesIO(PlanXlsxExporter.build(candidate)), read_only=True)
            self.assertTrue(
                any(
                    row[1] == urgent.urgent_request.external_id
                    for row in book["Заявки"].iter_rows(min_row=2, values_only=True)
                )
            )
            book.close()
            await decisions.reject(pending.plan.id)
            current = await decisions.get_current(Region.VOSTOK, now.date())
            self.assertEqual(current.id, base_id)
            rejected_event = await session.get(ReplanningEvent, pending.event_id)
            self.assertEqual(rejected_event.approval_status, ApprovalStatus.REJECTED)
            base_snapshot = await self.planning(session).prepare_replan_base(Region.VOSTOK)
            self.assertFalse(base_snapshot.approved_events)
            report = await DailyReportService(UnitOfWork(session)).build_snapshot(now.date())
            self.assertFalse(report.regions[0].events)

        async with self.sessions() as session:
            approved_urgent = await create_event(urgent, self.planning(session))
        async with self.sessions() as session:
            await PlanService(UnitOfWork(session)).approve(approved_urgent.plan.id)
        async with self.sessions() as session:
            base_snapshot = await self.planning(session).prepare_replan_base(Region.VOSTOK)
            self.assertEqual(len(base_snapshot.approved_events), 1)
            self.assertEqual(
                base_snapshot.approved_events[0].request_id, approved_urgent.request_id
            )
            report = await DailyReportService(UnitOfWork(session)).build_snapshot(now.date())
            self.assertEqual(
                [item.id for item in report.regions[0].events], [approved_urgent.event_id]
            )
        async with self.sessions() as session:
            with self.assertRaises(PlanningUrgentRequestExists):
                await create_event(urgent, self.planning(session))

        connection = EventPlanningRequest(
            region=Region.VOSTOK,
            event_type=ReplanningEventType.URGENT_REQUEST,
            urgent_request=UrgentRequestPayload.model_validate(
                {
                    **urgent.urgent_request.model_dump(),
                    "external_id": 987654322,
                    "type_bk": RequestTypeBk.CONNECTION,
                    "type_hd": RequestTypeHd.CONNECTION_REQUEST,
                    "norm_minutes": 90,
                    "norm_minutes_without_travel": 70,
                    "priority": 2,
                    "required_skill": Skill.CONNECTION_AND_ORDERS,
                }
            ),
        )
        async with self.sessions() as session:
            connection_plan = await create_event(connection, self.planning(session))
            self.assertEqual(connection_plan.plan.kind, PlanKind.EVENT_REPLAN)
        async with self.sessions() as session:
            await PlanService(UnitOfWork(session)).reject(connection_plan.plan.id)

        async with self.sessions() as session:
            with self.assertRaises(PlanningCurrentPlanMissing):
                await create_event(
                    EventPlanningRequest(
                        region=Region.YUGOTSENTR,
                        event_type=ReplanningEventType.REQUEST_CANCELLED,
                        request_id=request_id,
                    ),
                    self.planning(session),
                )
        async with self.sessions() as session:
            with self.assertRaises(PlanningEventTargetMissing):
                await create_event(
                    EventPlanningRequest(
                        region=Region.VOSTOK,
                        event_type=ReplanningEventType.REQUEST_CANCELLED,
                        request_id=uuid.uuid7(),
                    ),
                    self.planning(session),
                )
        async with self.sessions() as session:
            active_request = await session.get(Request, request_id)
            active_request.status = RequestStatus.ON_THE_WAY
            await session.commit()
        async with self.sessions() as session:
            cancelled = await create_event(
                EventPlanningRequest(
                    region=Region.VOSTOK,
                    event_type=ReplanningEventType.REQUEST_CANCELLED,
                    request_id=request_id,
                ),
                self.planning(session),
            )
            self.assertEqual(
                (await session.get(Request, request_id)).status, RequestStatus.ON_THE_WAY
            )
        async with self.sessions() as session:
            await PlanService(UnitOfWork(session)).approve(cancelled.plan.id)
            self.assertEqual(
                (await session.get(Request, request_id)).status, RequestStatus.CANCELLED
            )
        async with self.sessions() as session:
            with self.assertRaises(PlanningRequestAlreadyCancelled):
                await create_event(
                    EventPlanningRequest(
                        region=Region.VOSTOK,
                        event_type=ReplanningEventType.REQUEST_CANCELLED,
                        request_id=request_id,
                    ),
                    self.planning(session),
                )

        async with self.sessions() as session:
            unavailable = await create_event(
                EventPlanningRequest(
                    region=Region.VOSTOK,
                    event_type=ReplanningEventType.ENGINEER_UNAVAILABLE,
                    engineer_id=engineer_id,
                ),
                self.planning(session),
            )
            self.assertTrue((await session.get(Engineer, engineer_id)).is_available)
        async with self.sessions() as session:
            engineer = await session.get(Engineer, engineer_id)
            engineer.region = Region.YUGOTSENTR
            await session.commit()
        async with self.sessions() as session:
            with self.assertRaises(PlanStateChangedError):
                await PlanService(UnitOfWork(session)).approve(unavailable.plan.id)
        async with self.sessions() as session:
            candidate = await session.get(Plan, unavailable.plan.id)
            event = await session.get(ReplanningEvent, unavailable.event_id)
            self.assertEqual(candidate.approval_status, ApprovalStatus.PENDING)
            self.assertEqual(event.approval_status, ApprovalStatus.PENDING)
            engineer = await session.get(Engineer, engineer_id)
            engineer.region = Region.VOSTOK
            await session.commit()
        async with self.sessions() as session:
            await PlanService(UnitOfWork(session)).approve(unavailable.plan.id)
            self.assertFalse((await session.get(Engineer, engineer_id)).is_available)
        async with self.sessions() as session:
            with self.assertRaises(PlanningEngineerStateConflict):
                await create_event(
                    EventPlanningRequest(
                        region=Region.VOSTOK,
                        event_type=ReplanningEventType.ENGINEER_UNAVAILABLE,
                        engineer_id=engineer_id,
                    ),
                    self.planning(session),
                )

        async with self.sessions() as session:
            available = await create_event(
                EventPlanningRequest(
                    region=Region.VOSTOK,
                    event_type=ReplanningEventType.ENGINEER_AVAILABLE,
                    engineer_id=engineer_id,
                ),
                self.planning(session),
            )
        async with self.sessions() as session:
            await PlanService(UnitOfWork(session)).approve(available.plan.id)
            self.assertTrue((await session.get(Engineer, engineer_id)).is_available)
        async with self.sessions() as session:
            with self.assertRaises(PlanningEngineerStateConflict):
                await create_event(
                    EventPlanningRequest(
                        region=Region.VOSTOK,
                        event_type=ReplanningEventType.ENGINEER_AVAILABLE,
                        engineer_id=engineer_id,
                    ),
                    self.planning(session),
                )

        async with self.sessions() as session:
            unavailable_again = await create_event(
                EventPlanningRequest(
                    region=Region.VOSTOK,
                    event_type=ReplanningEventType.ENGINEER_UNAVAILABLE,
                    engineer_id=engineer_id,
                ),
                self.planning(session),
            )
        async with self.sessions() as session:
            await PlanService(UnitOfWork(session)).approve(unavailable_again.plan.id)
            self.assertFalse((await session.get(Engineer, engineer_id)).is_available)

        async with self.sessions() as session:
            ordinary = await self.planning(session).replan([Region.VOSTOK], None, None)
            self.assertEqual(ordinary.status, "success")
            ordinary_id = ordinary.regions[0].plan_summary.id
        async with self.sessions() as session:
            stale_event = await create_event(
                EventPlanningRequest(
                    region=Region.VOSTOK,
                    event_type=ReplanningEventType.ENGINEER_AVAILABLE,
                    engineer_id=engineer_id,
                ),
                self.planning(session),
            )
        async with self.sessions() as session:
            await PlanService(UnitOfWork(session)).approve(ordinary_id)
        async with self.sessions() as session:
            with self.assertRaises(PlanNotPendingError):
                await PlanService(UnitOfWork(session)).approve(stale_event.plan.id)
            stale_plan = await session.get(Plan, stale_event.plan.id)
            stale_fact = await session.get(ReplanningEvent, stale_event.event_id)
            self.assertEqual(stale_plan.approval_status, ApprovalStatus.REJECTED)
            self.assertEqual(stale_fact.approval_status, ApprovalStatus.REJECTED)


class EventRequestValidationTest(unittest.TestCase):
    def test_one_target_and_server_owned_time(self) -> None:
        with self.assertRaises(ValidationError):
            EventPlanningRequest.model_validate(
                {
                    "region": Region.VOSTOK,
                    "event_type": ReplanningEventType.REQUEST_CANCELLED,
                    "request_id": str(uuid.uuid7()),
                    "engineer_id": str(uuid.uuid7()),
                }
            )
        with self.assertRaises(ValidationError):
            EventPlanningRequest.model_validate(
                {
                    "region": Region.VOSTOK,
                    "event_type": ReplanningEventType.REQUEST_CANCELLED,
                    "request_id": str(uuid.uuid7()),
                    "occurred_at": datetime.now(ZoneInfo("Europe/Moscow")).isoformat(),
                }
            )

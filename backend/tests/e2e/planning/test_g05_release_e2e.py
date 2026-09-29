"""G05 full-day integration against a fresh disposable g05_test_* PostgreSQL database."""

import os
import unittest
from datetime import UTC, datetime, timedelta
from io import BytesIO
from unittest.mock import AsyncMock, Mock
from zipfile import ZipFile
from zoneinfo import ZoneInfo

from fastapi import UploadFile
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.engine import make_url
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from src.api.planning.router import create_event, import_initial_planning_data
from src.api.planning.schemas import EventPlanningRequest, UrgentRequestPayload
from src.api.planning.service import PlanningService
from src.api.plans.export import PlanXlsxExporter
from src.api.plans.service import PlanService
from src.api.reports.service import DailyReportExportService, DailyReportService
from src.core.algorithm import AlgorithmService
from src.core.db.enums import Region, ReplanningEventType, RequestTypeBk, RequestTypeHd, Skill
from src.core.db.models import Plan
from src.core.db.uow import UnitOfWork
from src.core.s3 import ExportDownload, ExportKind

from tests.support.planning import FakeGeocoder, FakeMatrixService, FakeStorage, pair

DATABASE_URL = os.environ.get("G05_TEST_DATABASE_URL")


@unittest.skipUnless(DATABASE_URL, "G05_TEST_DATABASE_URL is not set")
class ReleaseEndToEndTest(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self) -> None:
        assert DATABASE_URL is not None
        url = make_url(DATABASE_URL)
        if (
            url.drivername != "postgresql+asyncpg"
            or url.host not in {"127.0.0.1", "localhost", "postgres"}
            or not (url.database or "").startswith("g05_test_")
        ):
            self.fail("G05 requires a disposable local g05_test_* database")
        self.engine = create_async_engine(
            DATABASE_URL, connect_args={"server_settings": {"timezone": "UTC"}}
        )
        self.sessions = async_sessionmaker(self.engine, expire_on_commit=False)
        self.storage = FakeStorage()
        async with self.sessions() as session:
            if await session.scalar(select(func.count()).select_from(Plan)):
                self.fail("G05 database must be freshly migrated and empty")

    async def asyncTearDown(self) -> None:
        await self.engine.dispose()

    def planning(self, session) -> PlanningService:
        return PlanningService(
            UnitOfWork(session),
            FakeGeocoder(),
            self.storage,
            AlgorithmService(),
            FakeMatrixService(),
        )

    async def test_approved_day_flows_through_plans_and_reports(self) -> None:
        now = datetime.now(ZoneInfo("Europe/Moscow")).replace(tzinfo=None)
        if now.hour == 23:
            self.skipTest("Initial approval needs enough of the working day left")
        files = pair("Восток") + pair("Югоцентр")
        uploads = [UploadFile(file=BytesIO(item.data), filename=item.filename) for item in files]
        async with self.sessions() as session:
            initial = await import_initial_planning_data(self.planning(session), uploads)
        self.assertEqual(initial.status, "success")
        initial_ids = {item.region: item.plan_summary.id for item in initial.regions}
        async with self.sessions() as session:
            decisions = PlanService(UnitOfWork(session))
            for plan_id in initial_ids.values():
                await decisions.approve(plan_id)

        async with self.sessions() as session:
            replan = await self.planning(session).replan([Region.VOSTOK], None, None)
        self.assertEqual(replan.status, "success")
        replan_id = replan.regions[0].plan_summary.id
        async with self.sessions() as session:
            await PlanService(UnitOfWork(session)).approve(replan_id)

        urgent_id = 987654600
        event_request = EventPlanningRequest(
            region=Region.VOSTOK,
            event_type=ReplanningEventType.URGENT_REQUEST,
            urgent_request=UrgentRequestPayload(
                external_id=urgent_id,
                type_bk=RequestTypeBk.GLOBAL_PROBLEM,
                type_hd=RequestTypeHd.EMERGENCY,
                district="Тестовый район",
                address="Тестовый адрес",
                is_gigabit=False,
                window_start=now + timedelta(minutes=20),
                window_end=datetime.combine(now.date(), datetime.max.time()).replace(microsecond=0),
                norm_minutes=100,
                norm_minutes_without_travel=80,
                priority=1,
                required_skill=Skill.EMERGENCY_WORKS,
            ),
        )
        async with self.sessions() as session:
            event = await create_event(event_request, self.planning(session))
        async with self.sessions() as session:
            await PlanService(UnitOfWork(session)).approve(event.plan.id)

        async with self.sessions() as session:
            decisions = PlanService(UnitOfWork(session))
            current = await decisions.get_current(Region.VOSTOK, now.date())
            detail = await decisions.get_by_id(event.plan.id)
            plan_snapshot = await decisions.get_snapshot_by_id(event.plan.id)
            reports = DailyReportService(UnitOfWork(session))
            report = await reports.build_snapshot(now.date())
            self.assertEqual(current.id, event.plan.id)
            self.assertEqual(detail.diff.base_plan_id, replan_id)
            self.assertTrue(
                any(item.request_id == event.request_id for item in detail.diff.requests)
            )
            self.assertEqual(len(report.regions), 2)
            east = next(item for item in report.regions if item.region == Region.VOSTOK)
            self.assertEqual([item.id for item in east.events], [event.event_id])
            self.assertEqual(east.current_plan_id, event.plan.id)

            book = load_workbook(BytesIO(PlanXlsxExporter.build(plan_snapshot)), read_only=True)
            self.assertTrue(
                any(
                    row[1] == urgent_id
                    for row in book["Заявки"].iter_rows(min_row=2, values_only=True)
                )
            )
            book.close()

            delivery = Mock()
            delivery.deliver = AsyncMock(
                return_value=ExportDownload(
                    url="https://example.invalid/report",
                    expires_at=datetime.now(UTC) + timedelta(minutes=15),
                    filename=f"daily-report-{now.date().isoformat()}.zip",
                    content_type="application/zip",
                    size_bytes=1,
                )
            )
            download = await DailyReportExportService(reports, delivery).export(now.date())
            self.assertEqual(download.url, "https://example.invalid/report")
            args = delivery.deliver.await_args.kwargs
            self.assertEqual(args["kind"], ExportKind.DAILY_REPORT)
            with ZipFile(BytesIO(args["data"])) as archive:
                self.assertEqual(len(archive.namelist()), 3)
                self.assertTrue(
                    all(archive.read(name).startswith(b"%PDF-") for name in archive.namelist())
                )

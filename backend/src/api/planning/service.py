import asyncio
import logging
import uuid
from collections import defaultdict
from dataclasses import replace
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from src.api.planning.dto import (
    InitialPlanningResult,
    ParsedEngineer,
    ParsedRequest,
    ParsedWorkbook,
    PlanningImportResult,
    PlanningUploadFile,
)
from src.api.planning.service_exc import (
    PlanningFileCountError,
    PlanningFileValidationError,
    PlanningRegionPairError,
    RepeatedRequestError,
)
from src.api.planning.utils import parse_datetime, parse_type_bk, request_rules, required_string
from src.config import cfg
from src.core.algorithm import AlgorithmService, DistributionMode
from src.core.db.dto import EngineerCreateDTO, RequestCreateDTO, UploadedFileCreateDTO
from src.core.db.enums import ConnectionType, Region, RequestTypeHd, Skill, VehicleType
from src.core.db.uow import UnitOfWork
from src.core.geocoding import Coordinates, GeocodingService
from src.core.s3 import S3Storage

log = logging.getLogger(__name__)


class PlanningService:
    """Импортирует исходные таблицы для первичного планирования."""

    def __init__(
        self,
        uow: UnitOfWork,
        geocoding: GeocodingService,
        storage: S3Storage,
        algorithm: AlgorithmService,
    ) -> None:
        self._uow = uow
        self._geocoding = geocoding
        self._storage = storage
        self._algorithm = algorithm

    async def import_initial_data(self, files: list[PlanningUploadFile]) -> InitialPlanningResult:
        if not files or len(files) % 2 != 0:
            raise PlanningFileCountError

        parsed_workbooks = await asyncio.gather(
            *(asyncio.to_thread(self._parse_workbook, file) for file in files)
        )
        region_pairs = self._build_region_pairs(parsed_workbooks)
        all_external_ids = [
            request.external_id
            for pair in region_pairs.values()
            for request in pair["requests"].requests
        ]
        if len(all_external_ids) != len(set(all_external_ids)):
            raise RepeatedRequestError

        existing_ids = await self._uow.requests.get_existing_external_ids(set(all_external_ids))
        if existing_ids:
            raise RepeatedRequestError

        prepared_regions = []
        for region in sorted(region_pairs, key=lambda item: item.value):
            requests_workbook = region_pairs[region]["requests"]
            engineers_workbook = region_pairs[region]["engineers"]
            prepared_regions.append(
                await self._prepare_region(region, requests_workbook, engineers_workbook)
            )

        uploaded_objects: list[tuple[str, str]] = []
        results: list[PlanningImportResult] = []
        try:
            for prepared in prepared_regions:
                results.append(await self._persist_region(*prepared, uploaded_objects))
            await self._uow.commit()
        except Exception:
            await self._uow.rollback()
            await self._delete_uploaded_objects(uploaded_objects)
            raise

        plans = await asyncio.gather(
            *(
                self._algorithm.plan_initial(
                    result.upload_id, result.region, DistributionMode.MIN_ENGINEERS
                )
                for result in results
            )
        )
        results_with_plans = [
            replace(result, plan_id=plan.plan_id)
            for result, plan in zip(results, plans, strict=True)
        ]
        return InitialPlanningResult(status="imported", imports=tuple(results_with_plans))

    def _parse_workbook(self, source: PlanningUploadFile) -> ParsedWorkbook:
        if not source.filename or Path(source.filename).suffix.casefold() != ".xlsx":
            raise PlanningFileValidationError
        if not source.data or len(source.data) > cfg.planning.max_file_size_bytes:
            raise PlanningFileValidationError

        workbook = None
        try:
            workbook = load_workbook(BytesIO(source.data), read_only=True, data_only=True)
            worksheet = workbook.active
            title = required_string(worksheet.cell(row=1, column=1).value)
            region = self._parse_region(title)
            headers = [
                required_string(cell.value) if cell.value is not None else ""
                for cell in worksheet[2]
            ]
            header_indexes = {header: index for index, header in enumerate(headers) if header}
            role = "engineers" if "Бригада" in header_indexes else "requests"
            self._validate_workbook_kind(title, role)
            self._validate_headers(header_indexes, role)

            office_address = None
            requests: list[ParsedRequest] = []
            engineer_skills: dict[str, set[Skill]] = defaultdict(set)
            for row in worksheet.iter_rows(min_row=3, values_only=True):
                first_value = row[0]
                if isinstance(first_value, str) and first_value.casefold().startswith("адрес офис"):
                    office_address = required_string(row[1])
                    continue
                if first_value is None:
                    continue
                if role == "requests":
                    requests.append(self._parse_request(row, header_indexes))
                else:
                    engineer_name = row[header_indexes["Бригада"]]
                    if engineer_name is not None:
                        type_bk = parse_type_bk(row[header_indexes["Тип заявки BK"]])
                        engineer_skills[required_string(engineer_name)].add(
                            request_rules(type_bk)[3]
                        )

            if role == "requests" and (not requests or office_address is None):
                raise PlanningFileValidationError
            if role == "engineers" and not engineer_skills:
                raise PlanningFileValidationError

            engineers = tuple(
                ParsedEngineer(
                    name=name, skills=tuple(sorted(skills, key=lambda skill: skill.value))
                )
                for name, skills in sorted(engineer_skills.items())
            )
            return ParsedWorkbook(
                source=source,
                region=region,
                role=role,
                office_address=office_address,
                requests=tuple(requests),
                engineers=engineers,
            )
        except PlanningFileValidationError:
            raise
        except (BadZipFile, InvalidFileException, OSError, TypeError, ValueError, KeyError) as exc:
            raise PlanningFileValidationError from exc
        finally:
            if workbook is not None:
                workbook.close()

    def _build_region_pairs(
        self, workbooks: list[ParsedWorkbook]
    ) -> dict[Region, dict[str, ParsedWorkbook]]:
        region_pairs: dict[Region, dict[str, ParsedWorkbook]] = defaultdict(dict)
        for workbook in workbooks:
            if workbook.role in region_pairs[workbook.region]:
                raise PlanningRegionPairError
            region_pairs[workbook.region][workbook.role] = workbook
        if any(set(pair) != {"requests", "engineers"} for pair in region_pairs.values()):
            raise PlanningRegionPairError
        return region_pairs

    async def _prepare_region(
        self,
        region: Region,
        requests_workbook: ParsedWorkbook,
        engineers_workbook: ParsedWorkbook,
    ) -> tuple[
        Region,
        ParsedWorkbook,
        ParsedWorkbook,
        dict[str, Coordinates],
    ]:
        if requests_workbook.office_address is None:
            raise PlanningFileValidationError
        addresses = [requests_workbook.office_address]
        addresses.extend(request.address for request in requests_workbook.requests)
        coordinates = {}
        for address in dict.fromkeys(addresses):
            point = await self._geocoding.geocode(address)
            coordinates[address] = point
        return region, requests_workbook, engineers_workbook, coordinates

    async def _persist_region(
        self,
        region: Region,
        requests_workbook: ParsedWorkbook,
        engineers_workbook: ParsedWorkbook,
        coordinates: dict[str, Coordinates],
        uploaded_objects: list[tuple[str, str]],
    ) -> PlanningImportResult:
        upload_id = self._uow.data_uploads.create(region)
        file_dtos: list[UploadedFileCreateDTO] = []
        for workbook in (requests_workbook, engineers_workbook):
            bucket = cfg.s3.bucket_uploads
            safe_filename = Path(workbook.source.filename).name
            key = f"planning/{upload_id}/{uuid.uuid7()}-{safe_filename}"
            await self._storage.upload_file(
                bucket,
                key,
                workbook.source.data,
                workbook.source.content_type,
            )
            uploaded_objects.append((bucket, key))
            file_dtos.append(
                UploadedFileCreateDTO(
                    upload_id=upload_id,
                    filename=safe_filename,
                    content_type=workbook.source.content_type,
                    size_bytes=len(workbook.source.data),
                    s3_bucket=bucket,
                    s3_key=key,
                )
            )

        office_address = requests_workbook.office_address
        if office_address is None:
            raise PlanningFileValidationError
        office_coordinates = coordinates[office_address]
        planning_date = min(request.window_start for request in requests_workbook.requests).date()
        vehicle_type = VehicleType(cfg.planning.default_vehicle_type)

        request_dtos = [
            RequestCreateDTO(
                upload_id=upload_id,
                external_id=request.external_id,
                type_bk=request.type_bk,
                type_hd=request.type_hd,
                region=region,
                district=request.district,
                address=request.address,
                latitude=coordinates[request.address].latitude,
                longitude=coordinates[request.address].longitude,
                connection_type=request.connection_type,
                is_gigabit=request.is_gigabit,
                window_start=request.window_start,
                window_end=request.window_end,
                norm_minutes=request.norm_minutes,
                norm_minutes_without_travel=request.norm_minutes_without_travel,
                priority=request.priority,
                required_skill=request.required_skill,
                required_vehicle_type=None,
            )
            for request in requests_workbook.requests
        ]
        engineer_dtos = [
            EngineerCreateDTO(
                upload_id=upload_id,
                name=engineer.name,
                region=region,
                start_point_address=office_address,
                start_point_latitude=office_coordinates.latitude,
                start_point_longitude=office_coordinates.longitude,
                shift_start=datetime.combine(planning_date, cfg.planning.default_shift_start),
                shift_end=datetime.combine(planning_date, cfg.planning.default_shift_end),
                skills=engineer.skills,
                vehicle_type=vehicle_type,
            )
            for engineer in engineers_workbook.engineers
        ]
        self._uow.uploaded_files.add_many(file_dtos)
        self._uow.requests.add_many(request_dtos)
        self._uow.engineers.add_many(engineer_dtos)
        return PlanningImportResult(
            upload_id=upload_id,
            region=region,
            requests_count=len(request_dtos),
            engineers_count=len(engineer_dtos),
        )

    async def _delete_uploaded_objects(self, objects: list[tuple[str, str]]) -> None:
        for bucket, key in objects:
            try:
                await self._storage.delete_file(bucket, key)
            except Exception as exc:
                log.error(
                    "Failed to clean up uploaded planning file",
                    exc_info=exc,
                    extra={"s3": {"bucket": bucket, "key": key}},
                )

    @staticmethod
    def _parse_request(row: tuple[object, ...], indexes: dict[str, int]) -> ParsedRequest:
        type_bk = parse_type_bk(row[indexes["Тип заявки BK"]])
        norm, norm_without_travel, priority, skill = request_rules(type_bk)
        window_start = parse_datetime(row[indexes["Начало"]])
        window_end = parse_datetime(row[indexes["Окончание"]])
        if window_start >= window_end:
            raise PlanningFileValidationError

        connection_value = row[indexes["Подключение"]] if "Подключение" in indexes else None
        return ParsedRequest(
            external_id=PlanningService._parse_external_id(row[indexes["Заявка"]]),
            type_bk=type_bk,
            type_hd=PlanningService._parse_type_hd(row[indexes["Тип заявки HD"]]),
            district=required_string(row[indexes["Район"]]),
            address=required_string(row[indexes["Адрес"]]),
            connection_type=PlanningService._parse_connection_type(connection_value),
            is_gigabit=required_string(row[indexes["Гигабитное подключение"]]).casefold() == "да",
            window_start=window_start,
            window_end=window_end,
            norm_minutes=norm,
            norm_minutes_without_travel=norm_without_travel,
            priority=priority,
            required_skill=skill,
        )

    @staticmethod
    def _parse_region(title: str) -> Region:
        normalized = title.casefold()
        if "юго-восток" in normalized:
            return Region.YUGO_VOSTOK
        if "югоцентр" in normalized:
            return Region.YUGOTSENTR
        if "восток" in normalized:
            return Region.VOSTOK
        raise PlanningFileValidationError

    @staticmethod
    def _validate_workbook_kind(title: str, role: str) -> None:
        normalized = title.casefold()
        expected_marker = (
            "синтетические данные" if role == "requests" else "контрольное распределение"
        )
        if expected_marker not in normalized:
            raise PlanningFileValidationError

    @staticmethod
    def _validate_headers(indexes: dict[str, int], role: str) -> None:
        common = {"Заявка", "Тип заявки BK"}
        required = (
            common
            | {
                "Тип заявки HD",
                "Начало",
                "Окончание",
                "Район",
                "Адрес",
                "Гигабитное подключение",
            }
            if role == "requests"
            else common | {"Бригада"}
        )
        if not required.issubset(indexes):
            raise PlanningFileValidationError

    @staticmethod
    def _parse_type_hd(value: object) -> RequestTypeHd:
        mapping = {
            "IP-адрес 169...": RequestTypeHd.IP_ADDRESS_169,
            "TVE/ENT. Другие ошибки": RequestTypeHd.TVE_ENT_OTHER_ERRORS,
            "TVE/ENT. Замена приставки техником": (RequestTypeHd.TVE_ENT_SET_TOP_BOX_REPLACEMENT),
            "Авария": RequestTypeHd.EMERGENCY,
            "Дозаказ оборудования": RequestTypeHd.EQUIPMENT_ADDITIONAL_ORDER,
            "Заказ подключения/Дозаказ оборудования": (RequestTypeHd.CONNECTION_OR_EQUIPMENT_ORDER),
            "Заявка на подключение": RequestTypeHd.CONNECTION_REQUEST,
            "Информация": RequestTypeHd.INFORMATION,
            "Конвергенция абонента": RequestTypeHd.SUBSCRIBER_CONVERGENCE,
            "Мониторинг": RequestTypeHd.MONITORING,
            "Нет линка": RequestTypeHd.NO_LINK,
            "Низкая скорость": RequestTypeHd.LOW_SPEED,
            "Переключение на Гбит/с": RequestTypeHd.GIGABIT_SWITCH,
            "Работа с кабелем": RequestTypeHd.CABLE_WORK,
            "Разрывы": RequestTypeHd.DISCONNECTS,
            "Рост ошибок на порту": RequestTypeHd.PORT_ERROR_GROWTH,
            "Роутер. Замена техническим специалистом": (
                RequestTypeHd.ROUTER_REPLACEMENT_BY_TECHNICIAN
            ),
            "ТВ. Замена приставки техником": RequestTypeHd.TV_SET_TOP_BOX_REPLACEMENT,
        }
        try:
            return mapping[required_string(value)]
        except KeyError as exc:
            raise PlanningFileValidationError from exc

    @staticmethod
    def _parse_connection_type(value: object) -> ConnectionType | None:
        if value is None:
            return None
        try:
            return ConnectionType(required_string(value).casefold())
        except ValueError as exc:
            raise PlanningFileValidationError from exc

    @staticmethod
    def _parse_external_id(value: object) -> int:
        if isinstance(value, bool) or not isinstance(value, (int, float, str)):
            raise PlanningFileValidationError
        try:
            parsed = int(value)
        except (TypeError, ValueError) as exc:
            raise PlanningFileValidationError from exc
        if parsed <= 0 or isinstance(value, float) and not value.is_integer():
            raise PlanningFileValidationError
        return parsed

from collections import defaultdict
from datetime import datetime, time
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from src.api.exc.planning import (
    PlanningFileValidationError,
    PlanningRegionPairError,
)
from src.api.planning.dto import (
    ParsedEngineer,
    ParsedRequest,
    ParsedWorkbook,
    PlanningUploadFile,
)
from src.api.planning.utils import (
    parse_datetime,
    parse_type_bk,
    required_string,
    work_norm,
)
from src.config import cfg
from src.core.db.enums import ConnectionType, Region, RequestTypeHd, Skill, VehicleType


class PlanningWorkbookParser:
    @staticmethod
    def _validate_source(source: PlanningUploadFile) -> None:
        if not source.data or len(source.data) > cfg.planning.max_file_size_bytes:
            raise PlanningFileValidationError

    @staticmethod
    def identify_region(source: PlanningUploadFile) -> Region:
        """Читает только заголовок, чтобы ошибка книги затронула её округ."""
        PlanningWorkbookParser._validate_source(source)
        workbook = None
        try:
            workbook = load_workbook(BytesIO(source.data), read_only=True, data_only=True)
            title = required_string(workbook.active.cell(row=1, column=1).value)
            return PlanningWorkbookParser._parse_region(title)
        except PlanningFileValidationError:
            raise
        except (BadZipFile, InvalidFileException, OSError, TypeError, ValueError, KeyError) as exc:
            raise PlanningFileValidationError from exc
        finally:
            if workbook is not None:
                workbook.close()

    @staticmethod
    def parse(source: PlanningUploadFile) -> ParsedWorkbook:
        PlanningWorkbookParser._validate_source(source)

        workbook = None
        try:
            workbook = load_workbook(BytesIO(source.data), read_only=True, data_only=True)
            worksheet = workbook.active
            title = required_string(worksheet.cell(row=1, column=1).value)
            region = PlanningWorkbookParser._parse_region(title)
            headers = [
                required_string(cell.value) if cell.value is not None else ""
                for cell in worksheet[2]
            ]
            header_indexes = {header: index for index, header in enumerate(headers) if header}
            role = PlanningWorkbookParser._parse_role(header_indexes)

            office_address = None
            requests: list[ParsedRequest] = []
            engineers: list[ParsedEngineer] = []
            for row in worksheet.iter_rows(min_row=3, values_only=True):
                first_value = row[0]
                if (
                    role == "requests"
                    and isinstance(first_value, str)
                    and first_value.casefold().startswith("адрес офис")
                ):
                    office_address = required_string(row[1])
                    continue
                if first_value is None:
                    continue
                if role == "requests":
                    requests.append(PlanningWorkbookParser._parse_request(row, header_indexes))
                else:
                    engineers.append(PlanningWorkbookParser._parse_engineer(row, header_indexes))

            if role == "requests" and (not requests or office_address is None):
                raise PlanningFileValidationError
            if role == "engineers" and not engineers:
                raise PlanningFileValidationError
            engineer_names = [engineer.name for engineer in engineers]
            if len(engineer_names) != len(set(engineer_names)):
                raise PlanningFileValidationError
            return ParsedWorkbook(
                source=source,
                region=region,
                role=role,
                office_address=office_address,
                requests=tuple(requests),
                engineers=tuple(engineers),
            )
        except PlanningFileValidationError:
            raise
        except (BadZipFile, InvalidFileException, OSError, TypeError, ValueError, KeyError) as exc:
            raise PlanningFileValidationError from exc
        finally:
            if workbook is not None:
                workbook.close()

    @staticmethod
    def build_region_pairs(
        workbooks: list[ParsedWorkbook],
    ) -> dict[Region, dict[str, ParsedWorkbook]]:
        region_pairs: dict[Region, dict[str, ParsedWorkbook]] = defaultdict(dict)
        for workbook in workbooks:
            if workbook.role in region_pairs[workbook.region]:
                raise PlanningRegionPairError
            region_pairs[workbook.region][workbook.role] = workbook
        if any(set(pair) != {"requests", "engineers"} for pair in region_pairs.values()):
            raise PlanningRegionPairError
        return region_pairs

    @staticmethod
    def _parse_request(row: tuple[object, ...], indexes: dict[str, int]) -> ParsedRequest:
        type_bk = parse_type_bk(row[indexes["Тип заявки BK"]])
        norm = work_norm(type_bk)
        window_start = parse_datetime(row[indexes["Начало"]])
        window_end = parse_datetime(row[indexes["Окончание"]])
        if window_start >= window_end:
            raise PlanningFileValidationError

        connection_value = row[indexes["Подключение"]] if "Подключение" in indexes else None
        return ParsedRequest(
            external_id=PlanningWorkbookParser._parse_external_id(row[indexes["Заявка"]]),
            type_bk=type_bk,
            type_hd=PlanningWorkbookParser._parse_type_hd(row[indexes["Тип заявки HD"]]),
            district=required_string(row[indexes["Район"]]),
            address=required_string(row[indexes["Адрес"]]),
            connection_type=PlanningWorkbookParser._parse_connection_type(connection_value),
            is_gigabit=required_string(row[indexes["Гигабитное подключение"]]).casefold() == "да",
            window_start=window_start,
            window_end=window_end,
            norm_minutes=norm.minutes(with_travel=True),
            norm_minutes_without_travel=norm.minutes(with_travel=False),
            priority=norm.priority,
            required_skill=norm.required_skill,
        )

    @staticmethod
    def _parse_engineer(row: tuple[object, ...], indexes: dict[str, int]) -> ParsedEngineer:
        shift_start = PlanningWorkbookParser._parse_time(row[indexes["Начало смены"]])
        shift_end = PlanningWorkbookParser._parse_time(row[indexes["Конец смены"]])
        if shift_start >= shift_end:
            raise PlanningFileValidationError

        skill_columns = ("Навык 1", "Навык 2", "Навык 3")
        skills = tuple(
            PlanningWorkbookParser._parse_skill(row[indexes[column]])
            for column in skill_columns
            if column in indexes and row[indexes[column]] is not None
        )
        if not skills or len(skills) != len(set(skills)):
            raise PlanningFileValidationError

        return ParsedEngineer(
            name=required_string(row[indexes["Инженер"]]),
            start_point_address=required_string(row[indexes["Стартовая точка"]]),
            shift_start=shift_start,
            shift_end=shift_end,
            skills=skills,
            vehicle_type=PlanningWorkbookParser._parse_vehicle_type(row[indexes["Тип транспорта"]]),
        )

    @staticmethod
    def _parse_region(title: str) -> Region:
        normalized = title.casefold()
        if "юго-восток" in normalized:
            return Region.YUGO_VOSTOK
        if "югоцентр" in normalized:
            return Region.YUGOTSENTR
        if "восток" in normalized:
            return Region.VOSTOK
        raise PlanningFileValidationError

    @staticmethod
    def _parse_role(indexes: dict[str, int]) -> str:
        request_headers = {
            "Заявка",
            "Тип заявки BK",
            "Тип заявки HD",
            "Начало",
            "Окончание",
            "Район",
            "Адрес",
            "Гигабитное подключение",
        }
        engineer_headers = {
            "Инженер",
            "Стартовая точка",
            "Начало смены",
            "Конец смены",
            "Навык 1",
            "Тип транспорта",
        }
        matches = [
            role
            for role, required_headers in (
                ("requests", request_headers),
                ("engineers", engineer_headers),
            )
            if required_headers.issubset(indexes)
        ]
        if len(matches) != 1:
            raise PlanningFileValidationError
        return matches[0]

    @staticmethod
    def _parse_type_hd(value: object) -> RequestTypeHd:
        mapping = {
            "IP-адрес 169...": RequestTypeHd.IP_ADDRESS_169,
            "TVE/ENT. Другие ошибки": RequestTypeHd.TVE_ENT_OTHER_ERRORS,
            "TVE/ENT. Замена приставки техником": RequestTypeHd.TVE_ENT_SET_TOP_BOX_REPLACEMENT,
            "Авария": RequestTypeHd.EMERGENCY,
            "Дозаказ оборудования": RequestTypeHd.EQUIPMENT_ADDITIONAL_ORDER,
            "Заказ подключения/Дозаказ оборудования": (RequestTypeHd.CONNECTION_OR_EQUIPMENT_ORDER),
            "Заявка на подключение": RequestTypeHd.CONNECTION_REQUEST,
            "Информация": RequestTypeHd.INFORMATION,
            "Конвергенция абонента": RequestTypeHd.SUBSCRIBER_CONVERGENCE,
            "Мониторинг": RequestTypeHd.MONITORING,
            "Нет линка": RequestTypeHd.NO_LINK,
            "Низкая скорость": RequestTypeHd.LOW_SPEED,
            "Переключение на Гбит/с": RequestTypeHd.GIGABIT_SWITCH,
            "Работа с кабелем": RequestTypeHd.CABLE_WORK,
            "Разрывы": RequestTypeHd.DISCONNECTS,
            "Рост ошибок на порту": RequestTypeHd.PORT_ERROR_GROWTH,
            "Роутер. Замена техническим специалистом": (
                RequestTypeHd.ROUTER_REPLACEMENT_BY_TECHNICIAN
            ),
            "ТВ. Замена приставки техником": RequestTypeHd.TV_SET_TOP_BOX_REPLACEMENT,
        }
        try:
            return mapping[required_string(value)]
        except KeyError as exc:
            raise PlanningFileValidationError from exc

    @staticmethod
    def _parse_connection_type(value: object) -> ConnectionType | None:
        if value is None:
            return None
        try:
            return ConnectionType(required_string(value).casefold())
        except ValueError as exc:
            raise PlanningFileValidationError from exc

    @staticmethod
    def _parse_skill(value: object) -> Skill:
        mapping = {
            "Локальные работы": Skill.LOCAL_WORKS,
            "Работы на подключение и дозаказы": Skill.CONNECTION_AND_ORDERS,
            "Аварийные работы": Skill.EMERGENCY_WORKS,
        }
        try:
            return mapping[required_string(value)]
        except KeyError as exc:
            raise PlanningFileValidationError from exc

    @staticmethod
    def _parse_vehicle_type(value: object) -> VehicleType:
        mapping = {
            "Автомобиль": VehicleType.CAR,
            "Пешеход": VehicleType.PEDESTRIAN,
            "Велосипед": VehicleType.BICYCLE,
            "Общественный транспорт": VehicleType.PUBLIC_TRANSPORT,
        }
        try:
            return mapping[required_string(value)]
        except KeyError as exc:
            raise PlanningFileValidationError from exc

    @staticmethod
    def _parse_time(value: object) -> time:
        if isinstance(value, datetime):
            return value.time()
        if isinstance(value, time):
            return value
        try:
            return time.fromisoformat(required_string(value))
        except ValueError as exc:
            raise PlanningFileValidationError from exc

    @staticmethod
    def _parse_external_id(value: object) -> int:
        if isinstance(value, bool) or not isinstance(value, (int, float, str)):
            raise PlanningFileValidationError
        try:
            parsed = int(value)
        except (TypeError, ValueError) as exc:
            raise PlanningFileValidationError from exc
        if parsed <= 0 or isinstance(value, float) and not value.is_integer():
            raise PlanningFileValidationError
        return parsed

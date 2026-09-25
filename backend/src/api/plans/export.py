import io
import re
from datetime import UTC, date, datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.api.plans.diff_dto import PlanSnapshotDTO, SnapshotEngineerDTO, SnapshotRequestDTO


class PlanXlsxExporter:
    """Строит XLSX из уже загруженного неизменяемого снимка плана."""

    @staticmethod
    def build(snapshot: PlanSnapshotDTO) -> bytes:
        workbook = Workbook()
        metadata = workbook.active
        if metadata is None:
            raise ValueError("Workbook has no active worksheet")
        metadata.title = "План"
        PlanXlsxExporter._metadata(metadata, snapshot)

        requests = workbook.create_sheet("Заявки")
        PlanXlsxExporter._requests(requests, snapshot)

        engineers = workbook.create_sheet("Инженеры")
        PlanXlsxExporter._engineers(engineers, snapshot)

        unassigned = workbook.create_sheet("Неназначенные")
        PlanXlsxExporter._unassigned(unassigned, snapshot)

        metrics = workbook.create_sheet("Метрики")
        PlanXlsxExporter._metrics(metrics, snapshot)

        used_names = {sheet.title.casefold() for sheet in workbook.worksheets}
        for engineer in sorted(snapshot.engineers, key=lambda item: str(item.engineer_id)):
            sheet = workbook.create_sheet(PlanXlsxExporter._sheet_name(engineer.name, used_names))
            PlanXlsxExporter._engineer_route(sheet, engineer)

        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def _metadata(sheet: Worksheet, snapshot: PlanSnapshotDTO) -> None:
        metrics = snapshot.metrics
        rows: tuple[tuple[str, object], ...] = (
            ("ID плана", str(snapshot.id)),
            ("Округ", snapshot.region.value),
            ("Рабочая дата", snapshot.planning_date),
            ("Тип", snapshot.kind.value),
            ("Статус", snapshot.approval_status.value),
            ("Создан (МСК)", PlanXlsxExporter._moscow_time(snapshot.created_at, utc=True)),
            (
                "Утверждён (МСК)",
                PlanXlsxExporter._moscow_time(snapshot.approved_at, utc=True),
            ),
            (
                "Отклонён (МСК)",
                PlanXlsxExporter._moscow_time(snapshot.rejected_at, utc=True),
            ),
            ("Срез расчёта (МСК)", snapshot.calculation_cutoff_at),
            ("Базовый план", str(snapshot.based_on_plan_id) if snapshot.based_on_plan_id else None),
            (
                "Событие",
                str(snapshot.triggered_by_event_id) if snapshot.triggered_by_event_id else None,
            ),
            ("Назначено заявок", metrics.assigned_requests_count),
            ("Не назначено заявок", metrics.unassigned_requests_count),
            ("Задействовано инженеров", metrics.engineers_used_count),
            ("Доступно инженеров", metrics.available_engineers_count),
            ("Работа, мин", metrics.total_work_minutes),
            ("Дорога, мин", metrics.total_travel_minutes),
            ("Пробег, км", metrics.total_mileage_km),
            ("Средняя загрузка без дороги, %", metrics.average_workload_without_travel),
            ("Средняя загрузка с дорогой, %", metrics.average_workload_with_travel),
            (
                "Средняя загрузка занятых без дороги, %",
                metrics.average_used_workload_without_travel,
            ),
            (
                "Средняя загрузка занятых с дорогой, %",
                metrics.average_used_workload_with_travel,
            ),
            ("Минимальная загрузка с дорогой, %", metrics.min_workload_with_travel),
            ("Максимальная загрузка с дорогой, %", metrics.max_workload_with_travel),
        )
        sheet.append(("Показатель", "Значение"))
        for label, value in rows:
            sheet.append((label, PlanXlsxExporter._cell(value)))
        PlanXlsxExporter._finish_sheet(sheet)

    @staticmethod
    def _requests(sheet: Worksheet, snapshot: PlanSnapshotDTO) -> None:
        sheet.append(
            (
                "ID заявки",
                "Внешний ID",
                "Адрес",
                "Район",
                "Широта",
                "Долгота",
                "Окно с (МСК)",
                "Окно до (МСК)",
                "Приоритет",
                "Навык",
                "Норматив, мин",
                "Назначение",
                "ID инженера",
                "Порядок",
                "Прибытие (МСК)",
                "Начало (МСК)",
                "Конец (МСК)",
                "Дорога, мин",
                "Пробег, км",
                "Причина неназначения",
                "ID загрузки",
                "Тип BK",
                "Тип HD",
                "Подключение",
                "Гигабит",
                "Полный норматив, мин",
            )
        )
        for request in sorted(snapshot.requests, key=lambda item: str(item.request_id)):
            sheet.append(
                (
                    str(request.request_id),
                    request.external_id,
                    PlanXlsxExporter._cell(request.address),
                    PlanXlsxExporter._cell(request.district),
                    PlanXlsxExporter._cell(request.latitude),
                    PlanXlsxExporter._cell(request.longitude),
                    request.window_start,
                    request.window_end,
                    request.priority,
                    request.required_skill.value,
                    request.service_minutes,
                    "назначена" if request.engineer_id else "не назначена",
                    str(request.engineer_id) if request.engineer_id else None,
                    request.sequence_number,
                    request.planned_arrival,
                    request.planned_start,
                    request.planned_finish,
                    request.travel_minutes,
                    PlanXlsxExporter._cell(request.distance_km),
                    request.unassigned_reason.value if request.unassigned_reason else None,
                    str(request.upload_id) if request.upload_id else None,
                    request.type_bk.value if request.type_bk else None,
                    request.type_hd.value if request.type_hd else None,
                    request.connection_type.value if request.connection_type else None,
                    request.is_gigabit,
                    request.norm_minutes,
                )
            )
        PlanXlsxExporter._finish_sheet(sheet)

    @staticmethod
    def _engineers(sheet: Worksheet, snapshot: PlanSnapshotDTO) -> None:
        sheet.append(
            (
                "ID инженера",
                "Имя",
                "Транспорт",
                "Смена с (МСК)",
                "Смена до (МСК)",
                "Широта старта",
                "Долгота старта",
                "Доступен",
                "Заявок",
                "Работа, мин",
                "Дорога, мин",
                "Пробег, км",
                "Загрузка без дороги, %",
                "Загрузка с дорогой, %",
            )
        )
        for engineer in sorted(snapshot.engineers, key=lambda item: str(item.engineer_id)):
            sheet.append(
                (
                    str(engineer.engineer_id),
                    PlanXlsxExporter._cell(engineer.name),
                    engineer.vehicle_type.value,
                    engineer.shift_start,
                    engineer.shift_end,
                    PlanXlsxExporter._cell(engineer.start_latitude),
                    PlanXlsxExporter._cell(engineer.start_longitude),
                    engineer.is_available,
                    len(engineer.requests),
                    sum(item.service_minutes for item in engineer.requests),
                    sum(item.travel_minutes or 0 for item in engineer.requests),
                    PlanXlsxExporter._cell(engineer.route_distance_km),
                    PlanXlsxExporter._cell(engineer.workload_without_travel),
                    PlanXlsxExporter._cell(engineer.workload_with_travel),
                )
            )
        PlanXlsxExporter._finish_sheet(sheet)

    @staticmethod
    def _unassigned(sheet: Worksheet, snapshot: PlanSnapshotDTO) -> None:
        sheet.append(("ID заявки", "Внешний ID", "Адрес", "Приоритет", "Причина"))
        for request in sorted(snapshot.requests, key=lambda item: str(item.request_id)):
            if request.engineer_id is not None:
                continue
            sheet.append(
                (
                    str(request.request_id),
                    request.external_id,
                    PlanXlsxExporter._cell(request.address),
                    request.priority,
                    request.unassigned_reason.value if request.unassigned_reason else None,
                )
            )
        PlanXlsxExporter._finish_sheet(sheet)

    @staticmethod
    def _metrics(sheet: Worksheet, snapshot: PlanSnapshotDTO) -> None:
        metrics = snapshot.metrics
        sheet.append(("Показатель", "Значение"))
        rows = (
            ("Назначено заявок", metrics.assigned_requests_count),
            ("Не назначено заявок", metrics.unassigned_requests_count),
            ("Задействовано инженеров", metrics.engineers_used_count),
            ("Доступно инженеров", metrics.available_engineers_count),
            ("Работа, мин", metrics.total_work_minutes),
            ("Дорога, мин", metrics.total_travel_minutes),
            ("Пробег, км", metrics.total_mileage_km),
            ("Средняя загрузка без дороги, %", metrics.average_workload_without_travel),
            ("Средняя загрузка с дорогой, %", metrics.average_workload_with_travel),
            (
                "Средняя загрузка занятых без дороги, %",
                metrics.average_used_workload_without_travel,
            ),
            ("Средняя загрузка занятых с дорогой, %", metrics.average_used_workload_with_travel),
            ("Минимальная загрузка с дорогой, %", metrics.min_workload_with_travel),
            ("Максимальная загрузка с дорогой, %", metrics.max_workload_with_travel),
        )
        for label, value in rows:
            sheet.append((label, PlanXlsxExporter._cell(value)))
        PlanXlsxExporter._finish_sheet(sheet)

    @staticmethod
    def _engineer_route(sheet: Worksheet, engineer: SnapshotEngineerDTO) -> None:
        sheet.append(("Инженер", PlanXlsxExporter._cell(engineer.name)))
        sheet.append(("ID инженера", str(engineer.engineer_id)))
        sheet.append(
            (
                "Порядок",
                "ID заявки",
                "Внешний ID",
                "Адрес",
                "Прибытие (МСК)",
                "Начало (МСК)",
                "Конец (МСК)",
                "Работа, мин",
                "Дорога, мин",
                "Пробег, км",
                "Зафиксирована",
            )
        )
        for request in sorted(
            engineer.requests, key=lambda item: (item.sequence_number or 0, str(item.request_id))
        ):
            PlanXlsxExporter._route_row(sheet, request)
        PlanXlsxExporter._finish_sheet(sheet, header_row=3)

    @staticmethod
    def _route_row(sheet: Worksheet, request: SnapshotRequestDTO) -> None:
        sheet.append(
            (
                request.sequence_number,
                str(request.request_id),
                request.external_id,
                PlanXlsxExporter._cell(request.address),
                request.planned_arrival,
                request.planned_start,
                request.planned_finish,
                request.service_minutes,
                request.travel_minutes,
                PlanXlsxExporter._cell(request.distance_km),
                request.is_locked,
            )
        )

    @staticmethod
    def _sheet_name(name: str, used_names: set[str]) -> str:
        safe = re.sub(r"[\\/*?:\[\]\x00-\x1f]", "_", name).strip("' ") or "Инженер"
        candidate = safe[:31]
        suffix = 2
        while candidate.casefold() in used_names:
            tail = f"_{suffix}"
            candidate = f"{safe[: 31 - len(tail)]}{tail}"
            suffix += 1
        used_names.add(candidate.casefold())
        return candidate

    @staticmethod
    def _moscow_time(value: datetime | None, utc: bool) -> datetime | None:
        if value is None:
            return None
        if utc:
            source = value.replace(tzinfo=UTC) if value.tzinfo is None else value
            return source.astimezone(ZoneInfo("Europe/Moscow")).replace(tzinfo=None)
        return value

    @staticmethod
    def _cell(value: object) -> object:
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            return f"'{value}"
        if isinstance(value, Decimal):
            return float(value)
        return value

    @staticmethod
    def _finish_sheet(sheet: Worksheet, header_row: int = 1) -> None:
        sheet.freeze_panes = f"A{header_row + 1}"
        sheet.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        )
        for cell in sheet[header_row]:
            cell.font = Font(bold=True)
        for column in sheet.columns:
            cells = tuple(column)
            width = min(max(max(len(str(cell.value or "")) for cell in cells) + 2, 12), 42)
            sheet.column_dimensions[get_column_letter(cells[0].column)].width = width
            for cell in cells:
                if isinstance(cell.value, datetime):
                    cell.number_format = "dd.mm.yyyy hh:mm"
                elif isinstance(cell.value, date):
                    cell.number_format = "dd.mm.yyyy"

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import cast
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Flowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from beeline_backend.application.contracts import DayReport

BUSINESS_TIMEZONE = ZoneInfo("Europe/Moscow")


def _local_time(value: datetime) -> datetime:
    if value.tzinfo is None:
        return value
    return value.astimezone(BUSINESS_TIMEZONE)


def _excel_time(value: datetime | None) -> datetime | str:
    if value is None:
        return "нет данных"
    return _local_time(value).replace(tzinfo=None)


def _excel_text(value: str) -> str:
    if value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{value}"
    return value


def _excel_value(value: float | int | str) -> float | int | str:
    return _excel_text(value) if isinstance(value, str) else value


class XlsxReportRenderer:
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    extension = "xlsx"

    def render(self, report: DayReport, destination: Path) -> None:
        workbook = Workbook()
        sheet = cast(Worksheet, workbook.active)
        sheet.title = "Итог дня"
        sheet.append(["Итог дня", _excel_text(report.scenario)])
        sheet.append(["Дата", report.planning_date])
        sheet.append(["Версия отчёта", _excel_text(report.report_version)])
        sheet.append(["Статус", "Промежуточный" if report.interim else "Итоговый"])
        sheet.append([])
        headers = [
            "Инженер",
            "Позиция",
            "Заявка",
            "Адрес",
            "Плановое начало",
            "Плановое окончание",
            "Подтверждённый статус",
            "Фактическое начало",
            "Фактическое окончание",
            "Источник факта",
            "Плановое расстояние, м",
        ]
        sheet.append(headers)
        for item in report.assignments:
            sheet.append(
                [
                    _excel_text(item.engineer),
                    item.position,
                    _excel_text(item.request_external_id),
                    _excel_text(item.address),
                    _excel_time(item.planned_start),
                    _excel_time(item.planned_finish),
                    _excel_text(item.confirmed_status),
                    _excel_time(item.actual_start),
                    _excel_time(item.actual_finish),
                    _excel_text(item.fact_source) if item.fact_source else "нет данных",
                    item.distance_meters,
                ]
            )
        metric_sheet = workbook.create_sheet("Метрики")
        metric_sheet.append(["Метрика", "Значение"])
        for code, value in sorted(report.metrics.items()):
            metric_sheet.append([_excel_text(code), _excel_value(value)])
        header_fill = PatternFill("solid", fgColor="4A0072")
        for current in (sheet, metric_sheet):
            header_row = 6 if current is sheet else 1
            for cell in current[header_row]:
                cell.fill = header_fill
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            current.freeze_panes = f"A{header_row + 1}"
            for column in range(1, current.max_column + 1):
                width = min(
                    48,
                    max(
                        12,
                        max(
                            len(str(current.cell(row=row, column=column).value or ""))
                            for row in range(1, current.max_row + 1)
                        )
                        + 2,
                    ),
                )
                current.column_dimensions[get_column_letter(column)].width = width
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)


class PdfReportRenderer:
    media_type = "application/pdf"
    extension = "pdf"

    def __init__(self) -> None:
        self._font = "Helvetica"
        font_candidates = (
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
            Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
        )
        for font_path in font_candidates:
            if font_path.exists():
                pdfmetrics.registerFont(TTFont("BeelineReportFont", str(font_path)))
                self._font = "BeelineReportFont"
                break

    def render(self, report: DayReport, destination: Path) -> None:
        destination.parent.mkdir(parents=True, exist_ok=True)
        document = SimpleDocTemplate(
            str(destination),
            pagesize=landscape(A4),
            rightMargin=10 * mm,
            leftMargin=10 * mm,
            topMargin=10 * mm,
            bottomMargin=10 * mm,
        )
        styles = getSampleStyleSheet()
        for style in styles.byName.values():
            style.fontName = self._font
        story: list[Flowable] = [
            Paragraph(f"Итог дня: {report.scenario}", styles["Title"]),
            Paragraph(
                f"Дата: {report.planning_date.isoformat()}. Версия: {report.report_version}. "
                f"Статус: {'промежуточный' if report.interim else 'итоговый'}.",
                styles["BodyText"],
            ),
            Spacer(1, 5 * mm),
        ]
        data: list[list[str]] = [
            [
                "Инженер",
                "№",
                "Заявка",
                "Адрес",
                "План",
                "Статус",
                "Факт",
                "Источник",
                "План. км",
            ]
        ]
        for item in report.assignments:
            planned_start = _local_time(item.planned_start)
            planned_finish = _local_time(item.planned_finish)
            actual = (
                f"{_local_time(item.actual_start):%H:%M}-"
                f"{_local_time(item.actual_finish):%H:%M}"
                if item.actual_start and item.actual_finish
                else "нет данных"
            )
            data.append(
                [
                    item.engineer,
                    str(item.position),
                    item.request_external_id,
                    item.address,
                    f"{planned_start:%H:%M}-{planned_finish:%H:%M}",
                    item.confirmed_status,
                    actual,
                    item.fact_source or "нет данных",
                    f"{item.distance_meters / 1000:.1f}",
                ]
            )
        table = Table(
            data,
            repeatRows=1,
            colWidths=[34 * mm, 8 * mm, 18 * mm, 72 * mm, 25 * mm, 27 * mm, 27 * mm, 25 * mm, 18 * mm],
        )
        table.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), self._font),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4A0072")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("FONTSIZE", (0, 0), (-1, -1), 7),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F0F8")]),
                ]
            )
        )
        story.append(table)
        story.append(Spacer(1, 5 * mm))
        story.append(Paragraph("Метрики", styles["Heading2"]))
        metric_data = [[code, str(value)] for code, value in sorted(report.metrics.items())]
        metric_table = Table(metric_data, colWidths=[75 * mm, 40 * mm])
        metric_table.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), self._font),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                ]
            )
        )
        story.append(metric_table)
        document.build(story)

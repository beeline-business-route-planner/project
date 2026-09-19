from __future__ import annotations

import hashlib
import io
import re
from datetime import datetime
from pathlib import Path
from typing import cast
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from beeline_backend.application.contracts import (
    ImportedDataset,
    ImportedRequest,
    ImportIssue,
)
from beeline_backend.domain.errors import DomainError

SYNTHETIC_FILE_RE = re.compile(r"^(Восток|Югоцентр|Юго-восток) Синтетические данные\.xlsx$")
SERVICE_RULES: dict[str, tuple[str, int, int, str]] = {
    "connection": ("connection", 70, 90, "connection"),
    "emergency": ("emergency", 80, 100, "emergency"),
    "equipment_order": ("equipment_order", 20, 40, "connection"),
    "local": ("local", 30, 50, "local"),
}


def _parse_datetime(value: object, timezone: ZoneInfo) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, str):
        parsed = datetime.strptime(value.strip(), "%d.%m.%Y %H:%M")
    else:
        raise ValueError("expected date and time in DD.MM.YYYY HH:MM")
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone)
    return parsed


def _classify(bk_type: str, hd_type: str) -> tuple[str | None, int | None, int | None, str | None]:
    bk = bk_type.strip().casefold()
    hd = hd_type.strip().casefold()
    rule: str | None = None
    if bk == "дозаказ" or "дозаказ оборудования" in hd:
        rule = "equipment_order"
    elif bk == "подключение" and (
        "подключ" in hd or "конвергенция" in hd
    ):
        rule = "connection"
    elif bk == "глобальная проблема" and hd == "авария":
        rule = "emergency"
    elif bk == "локальная заявка":
        rule = "local"
    if rule is None:
        return None, None, None, None
    code, service, full, skill = SERVICE_RULES[rule]
    return code, service, full, skill


class XlsxDatasetImporter:
    def __init__(self, timezone: str = "Europe/Moscow") -> None:
        self._timezone = ZoneInfo(timezone)

    def parse(self, filename: str, content: bytes) -> ImportedDataset:
        clean_name = Path(filename).name
        match = SYNTHETIC_FILE_RE.fullmatch(clean_name)
        if not match or clean_name.startswith("~$") or "Контрольное" in clean_name:
            raise DomainError(
                "unsupported_dataset_file",
                "Only named synthetic XLSX files are accepted",
                {"file": clean_name, "allowed_pattern": SYNTHETIC_FILE_RE.pattern},
            )
        scenario_name = match.group(1)
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise DomainError("invalid_xlsx", "The uploaded XLSX cannot be read") from exc
        if len(workbook.sheetnames) != 1:
            raise DomainError("invalid_sheet_count", "Synthetic workbook must contain one sheet")
        sheet = cast(Worksheet, workbook.active)
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[2]]
        required = {"Заявка", "Тип заявки BK", "Тип заявки HD", "Начало", "Окончание", "Район", "Адрес"}
        missing = sorted(required - set(headers))
        if missing:
            raise DomainError(
                "missing_columns", "Synthetic workbook is missing required columns", {"columns": missing}
            )
        by_name = {name: index for index, name in enumerate(headers)}
        imported: list[ImportedRequest] = []
        issues: list[ImportIssue] = []
        office_address: str | None = None
        seen_external_ids: set[str] = set()
        for row_number, values in enumerate(sheet.iter_rows(min_row=3, values_only=True), start=3):
            first = values[0] if values else None
            if isinstance(first, str) and first.strip().casefold().startswith("адрес офис"):
                office_address = str(values[1]).strip() if len(values) > 1 and values[1] else None
                continue
            if first is None:
                continue
            row_errors: list[tuple[str, str, str]] = []
            try:
                external_id = str(int(first)) if isinstance(first, float) else str(first).strip()
            except (TypeError, ValueError):
                external_id = str(first).strip()
            if not external_id:
                row_errors.append(("Заявка", "missing_request_id", "Request ID is empty"))
            elif external_id in seen_external_ids:
                row_errors.append(("Заявка", "duplicate_request_id", "Request ID is duplicated"))
            bk_type = str(values[by_name["Тип заявки BK"]] or "").strip()
            hd_type = str(values[by_name["Тип заявки HD"]] or "").strip()
            district = str(values[by_name["Район"]] or "").strip()
            address = str(values[by_name["Адрес"]] or "").strip()
            if not address:
                row_errors.append(("Адрес", "missing_address", "Address is empty"))
            try:
                window_start = _parse_datetime(values[by_name["Начало"]], self._timezone)
                window_end = _parse_datetime(values[by_name["Окончание"]], self._timezone)
                if window_end <= window_start:
                    row_errors.append(("Окончание", "invalid_window", "Window end must be after start"))
            except ValueError as exc:
                row_errors.append(("Начало/Окончание", "invalid_datetime", str(exc)))
                window_start = datetime(1970, 1, 1, tzinfo=self._timezone)
                window_end = window_start
            for field, code, message in row_errors:
                issues.append(
                    ImportIssue(
                        file=clean_name,
                        sheet=sheet.title,
                        row=row_number,
                        field=field,
                        reason_code=code,
                        message=message,
                    )
                )
            if row_errors:
                continue
            seen_external_ids.add(external_id)
            work_code, service, full, skill = _classify(bk_type, hd_type)
            if work_code is None:
                issues.append(
                    ImportIssue(
                        file=clean_name,
                        sheet=sheet.title,
                        row=row_number,
                        field="Тип заявки BK/HD",
                        reason_code="needs_mapping",
                        message=f"No normative mapping for {bk_type!r} / {hd_type!r}",
                    )
                )
            gigabit_index = by_name.get("Гигабитное подключение")
            gigabit = (
                gigabit_index is not None
                and str(values[gigabit_index] or "").strip().casefold() == "да"
            )
            connection_index = by_name.get("Подключение")
            connection_kind = (
                str(values[connection_index]).strip()
                if connection_index is not None and values[connection_index]
                else None
            )
            imported.append(
                ImportedRequest(
                    row_number=row_number,
                    external_id=external_id,
                    bk_type=bk_type,
                    hd_type=hd_type,
                    window_start=window_start,
                    window_end=window_end,
                    district=district,
                    address=address,
                    gigabit=gigabit,
                    connection_kind=connection_kind,
                    work_code=work_code,
                    service_minutes=service,
                    full_normative_minutes=full,
                    required_skill=skill,
                    mapping_state="mapped" if work_code else "needs_mapping",
                )
            )
        if not imported:
            raise DomainError("empty_dataset", "Synthetic workbook contains no valid request rows")
        if office_address is None:
            raise DomainError("missing_office", "Synthetic workbook does not contain an office address")
        dates = {item.window_start.date() for item in imported}
        if len(dates) != 1:
            raise DomainError("mixed_planning_dates", "All requests must belong to one planning date")
        return ImportedDataset(
            scenario_code={"Восток": "east", "Югоцентр": "south-center", "Юго-восток": "south-east"}[scenario_name],
            scenario_name=scenario_name,
            office_address=office_address,
            planning_date=dates.pop(),
            source_name=clean_name,
            sha256=hashlib.sha256(content).hexdigest(),
            requests=imported,
            issues=issues,
        )
